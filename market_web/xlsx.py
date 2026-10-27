from __future__ import annotations

import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .data import FIELDS, connect


def link_cell(ws, value):
    """A native external-link relationship, never an Excel formula."""
    from urllib.parse import urlsplit
    cell = text_cell(ws, value)
    if isinstance(value, str):
        try:
            parsed = urlsplit(value)
            valid = parsed.scheme == 'https' and bool(parsed.hostname) and not parsed.username and not parsed.password
        except ValueError:
            valid = False
        if valid:
            cell.hyperlink = value
            cell.font = Font(color='2257D6', underline='single')
    return cell


def text_cell(ws, value):
    # Force text type even for formula-looking titles, keywords, shop names and long IDs.
    from openpyxl.cell import WriteOnlyCell
    cell = WriteOnlyCell(ws, value=value)
    if value is not None:
        cell.data_type = 's'
    return cell


def workbook(task):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('商品榜单')
    ws.freeze_panes = 'D2'
    ws.sheet_view.showGridLines = False
    for i, key in enumerate(FIELDS):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i + 1)].width = 48 if key in {'title', 'product_url', 'image_url'} else 22
    headers = []
    for label, _ in FIELDS.values():
        cell = text_cell(ws, label)
        cell.fill = PatternFill('solid', fgColor='2257D6')
        cell.font = Font(color='FFFFFF', bold=True)
        headers.append(cell)
    ws.append(headers)
    with connect() as c:
        for record in c.execute('SELECT data FROM items WHERE task_id=? ORDER BY position', (task['id'],)):
            row = json.loads(record[0])
            ws.append([(link_cell if k in {'product_url', 'image_url'} else text_cell)(ws, row.get(k)) for k in FIELDS])
    ws.auto_filter.ref = f'A1:K{task["count"] + 1}'
    notes = wb.create_sheet('采集说明')
    notes.column_dimensions['A'].width = 24
    notes.column_dimensions['B'].width = 100
    status = {'completed': '完成', 'failed': '失败 部分结果', 'cancelled': '已取消 部分结果', 'interrupted': '中断 部分结果'}.get(task['status'], task['status'])
    profile = task['profile']
    from .periods import LABELS, mode_of
    period = profile.get('period') or {}
    entries = [
        ('任务状态', status), ('任务编号', task['id']), ('创建时间 UTC', task['created_at']),
        ('更新时间 UTC', task['updated_at']), ('采集页数', str(task['pages'])), ('商品数', str(task['count'])),
        ('来源页面', profile.get('source_page')), ('来源路径', profile.get('path')),
        ('筛选条件', json.dumps(profile.get('params', {}), ensure_ascii=False)),
        ('结果说明', task['message']), ('未返回的字段', ' / '.join(task['missing']) or '无'),
        ('时间入口', period.get('label') or LABELS.get(mode_of(profile), '已识别周期')),
        ('统计范围', profile.get('params', {}).get('dateRange')),
        ('数据类型', period.get('kind') or ('实时单次快照' if mode_of(profile) == 'realtime' else '历史周期统计')),
        ('实时口径', '实时为逐页读取的单次采集 采集期间榜单可能变化 不代表同一瞬间的全量冻结结果' if mode_of(profile) == 'realtime' else '不适用'),
        ('指标口径', '区间保留平台原文 不代表精确成交人数或销量 缺失值留空'),
        ('商品主图', '首版保留图片链接 不内嵌下载图片'),
        ('无商品标识的记录', '若商品ID与链接均缺失 按原记录保留 不能保证跨页去重'),
        ('完成范围', '仅覆盖本任务筛选条件与账号可访问范围 不代表市场全量'),
    ]
    for key, value in entries:
        notes.append([text_cell(notes, key), (link_cell if key == '来源页面' else text_cell)(notes, value)])
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
