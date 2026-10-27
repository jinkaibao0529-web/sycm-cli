import io
import json
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from market_web.app import create_app
from market_web.data import connect, get_task, init, normalize, now, save_rows, update_task
from market_web.xlsx import workbook


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('MARKET_STATE_DIR', str(tmp_path))
    with TestClient(create_app()) as c:
        yield c


def seed(status='completed'):
    init()
    task_id = uuid.uuid4().hex
    profile = {'id':'capture', 'signature':'/mc/item/rank.json?cateId=1', 'path':'/mc/item/rank.json',
               'params':{'cateId':'1','dateRange':'2026-10-19|2026-10-25'},'source_page':'https://sycm.taobao.com/mc/rank',
               'account_key':'private-hash'}
    with connect() as c:
        c.execute('INSERT INTO tasks(id,status,created_at,updated_at,profile) VALUES (?,?,?,?,?)',
                  (task_id,status,now(),now(),json.dumps(profile)))
    rows = [normalize({'itemId':'12345678901234567890','itemTitle':'=HYPERLINK("https://invalid")','uv':{'value':'75万 ~ 100万'},'payByrCnt':{'value':'2.5万 ~ 5万'}})]
    count = save_rows(task_id, rows, 0)
    update_task(task_id,count=count,pages=1,missing='["店铺类型"]')
    return task_id


def test_normalize_ranges_and_nested_data():
    r = normalize({'item':{'itemId':123,'title':'商品'},'payByrCnt':{'value':'5000 ~ 7500'},'uv':{'range':'10万 ~ 25万'},'keywords':[{'value':'官方'},'超薄'],'picUrl':'//img.alicdn.com/a.jpg'})
    assert r['item_id']=='123' and r['title']=='商品'
    assert r['pay_buyers']=='5000 ~ 7500' and r['visitors']=='10万 ~ 25万'
    assert r['keywords']=='官方 / 超薄'
    assert r['image_url']=='https://img.alicdn.com/a.jpg'
    assert r['shop_name'] is None
    assert normalize({'itemUrl':'javascript:alert(1)'})['product_url'] is None


def test_frozen_windows_ipc_is_ascii_safe(capsys):
    from market_web.worker import emit
    emit({'message':'市场榜单已连接'})
    packet=capsys.readouterr().out
    assert packet.isascii()
    assert json.loads(packet)['message']=='市场榜单已连接'


def test_actual_market_endpoint_field_shape():
    r=normalize({'cateRankId':{'value':1,'cycleCqc':0},'coreKeyWord':{'value':'官方,正品'},
                 'itemId':{'value':123},'item':{'itemId':123,'title':'商品','detailUrl':'//item.taobao.com/item.htm?id=123','pictUrl':'//img.alicdn.com/a.jpg'},
                 'shop':{'title':'店铺','b2CShop':True},'payByrCnt':{'value':'2.5万 ~ 5万'},'uv':{'value':'75万 ~ 100万'}})
    assert all(v is not None for v in r.values())
    assert r['rank']=='1' and r['rank_change']=='0' and r['shop_type']=='天猫'


def test_rows_dedup_missing_ids_and_session_privacy(client):
    task_id = seed()
    row = normalize({'itemId':'12345678901234567890','title':'duplicate'})
    assert save_rows(task_id,[row],1)==1
    save_rows(task_id,[normalize({'title':'无ID'}),normalize({'title':'无ID'})],2)
    response=client.get('/api/tasks')
    assert 'private-hash' not in response.text and 'account_key' not in response.text
    assert client.get(f'/api/tasks/{task_id}/items?q=HYPERLINK').json()['total']==1
    assert client.get(f'/api/tasks/{task_id}/items?page=0').status_code==422
    assert client.get('/api/tasks/not-a-valid-id').status_code==404


def test_excel_text_safety_and_incomplete_provenance(client):
    task_id=seed('cancelled')
    response=client.get(f'/api/tasks/{task_id}/export')
    assert response.status_code==200
    wb=load_workbook(io.BytesIO(response.content))
    try:
        assert wb.sheetnames==['商品榜单','采集说明']
        row={cell.value: i+1 for i,cell in enumerate(wb['商品榜单'][1])}
        ws=wb['商品榜单']
        assert ws.cell(2,row['商品ID']).value=='12345678901234567890'
        assert ws.cell(2,row['商品名称']).data_type=='s'
        assert ws.cell(2,row['支付买家数区间']).value=='2.5万 ~ 5万'
        notes=dict(wb['采集说明'].iter_rows(values_only=True))
        assert '部分结果' in notes['任务状态']
        assert '店铺类型' in notes['未返回的字段']
    finally:
        wb.close()


def test_mutation_guard_and_invalid_source(client):
    assert client.post('/api/tasks',json={'profile_id':'fake'}).status_code==403
    token=client.get('/api/session').json()['token']
    headers={'X-Market-Token':token}
    assert client.post('/api/tasks',headers=headers,json={'profile_id':'fake'}).status_code==422
    assert client.post('/api/tasks',headers={**headers,'Origin':'https://evil.invalid'},json={'profile_id':'fake'}).status_code==403
    assert client.get('/api/health',headers={'Host':'evil.invalid'}).status_code==403
    assert client.get('/api/health').headers['X-Content-Type-Options']=='nosniff'


def test_cancel_running_task(client):
    task_id=seed('running')
    token=client.get('/api/session').json()['token']
    response=client.post(f'/api/tasks/{task_id}/cancel',headers={'X-Market-Token':token})
    assert response.json()['status']=='cancelling'
    assert client.get(f'/api/tasks/{task_id}/export').status_code==409


def test_restart_marks_running_task_interrupted(tmp_path,monkeypatch):
    monkeypatch.setenv('MARKET_STATE_DIR',str(tmp_path))
    task_id=seed('running')
    with TestClient(create_app()) as client:
        assert client.get(f'/api/tasks/{task_id}').json()['status']=='interrupted'


def test_account_change_invalidates_all_rows(client,monkeypatch):
    from market_web import worker
    task_id=seed('queued')
    class Connection:
        def __enter__(self):return self
        def __exit__(self,*_):pass
        def ask(self,**_):
            from tb.core.errors import AccountChanged
            raise AccountChanged('changed')
    monkeypatch.setattr(worker,'MarketConnection',Connection)
    worker.collect(task_id)
    t=get_task(task_id)
    assert t['status']=='failed' and t['count']==0
    assert client.get(f'/api/tasks/{task_id}/items').json()['total']==0


def test_repeated_page_stops_instead_of_false_success(client,monkeypatch):
    from market_web import worker
    task_id=seed('queued')
    class Connection:
        account_key='private-hash'
        def __enter__(self):return self
        def __exit__(self,*_):pass
        def ask(self,**_):return {'rows':[{'itemId':'123','itemTitle':'商品'}],'total':5,'page_key':'page'}
    monkeypatch.setattr(worker,'MarketConnection',Connection)
    monkeypatch.setattr(worker.random,'uniform',lambda *_:0)
    worker.collect(task_id)
    assert get_task(task_id)['status']=='failed'
    assert '重复页' in get_task(task_id)['message']


def test_pagination_completes_using_raw_row_count(client,monkeypatch):
    from market_web import worker
    task_id=seed('queued')
    with connect() as c:c.execute('DELETE FROM items WHERE task_id=?',(task_id,))
    class Connection:
        account_key='private-hash'
        def __enter__(self):return self
        def __exit__(self,*_):pass
        def ask(self,**args):return {'rows':[{'itemId':str(args['page']),'itemTitle':'商品'}],'total':2,'page_key':'page'}
    monkeypatch.setattr(worker,'MarketConnection',Connection)
    monkeypatch.setattr(worker.random,'uniform',lambda *_:0)
    worker.collect(task_id)
    assert get_task(task_id)['status']=='completed' and get_task(task_id)['count']==2
